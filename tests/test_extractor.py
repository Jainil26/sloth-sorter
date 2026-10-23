"""Tests for safe text extraction across supported and unsupported file formats."""

from pathlib import Path
import docx
import pymupdf
import pytest

from sloth_sorter.extractor import extract_text


def test_extract_txt(tmp_path):
    f = tmp_path / "sample.txt"
    f.write_text("Hello world! This is a simple text file.")
    extracted = extract_text(f)
    assert extracted == "Hello world! This is a simple text file."


def test_extract_md(tmp_path):
    f = tmp_path / "notes.md"
    f.write_text("# Title\n\nSome markdown notes.")
    extracted = extract_text(f)
    assert "# Title" in extracted
    assert "Some markdown notes." in extracted


def test_extract_csv(tmp_path):
    f = tmp_path / "data.csv"
    f.write_text("name,amount\ninvoice,120.50\n")
    extracted = extract_text(f)
    assert "invoice,120.50" in extracted


def test_extract_docx(tmp_path):
    f = tmp_path / "sample.docx"
    doc = docx.Document()
    doc.add_paragraph("First paragraph in docx.")
    doc.add_paragraph("Second paragraph with resume skills.")
    doc.save(str(f))

    extracted = extract_text(f)
    assert "First paragraph in docx." in extracted
    assert "Second paragraph with resume skills." in extracted


def test_extract_pdf(tmp_path):
    f = tmp_path / "sample.pdf"
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((50, 50), "Quarterly Financial Report Q3 2024")
    doc.save(str(f))
    doc.close()

    extracted = extract_text(f)
    assert "Quarterly Financial Report Q3 2024" in extracted


def test_extract_unsupported_formats(tmp_path):
    f_bin = tmp_path / "photo.jpg"
    f_bin.write_bytes(b"\xff\xd8\xff\xe0\x00\x10JFIF")
    assert extract_text(f_bin) == ""

    f_exe = tmp_path / "setup.exe"
    f_exe.write_bytes(b"MZ\x90\x00")
    assert extract_text(f_exe) == ""


def test_extract_corrupt_files_never_crashes(tmp_path):
    # Corrupt PDF
    f_pdf = tmp_path / "corrupt.pdf"
    f_pdf.write_bytes(b"%PDF-1.4 garbage binary data truncated")
    assert extract_text(f_pdf) == ""

    # Corrupt DOCX
    f_docx = tmp_path / "corrupt.docx"
    f_docx.write_bytes(b"not a valid zip or docx archive")
    assert extract_text(f_docx) == ""


def test_extract_cap_length(tmp_path):
    f = tmp_path / "long.txt"
    long_content = "A" * 5000
    f.write_text(long_content)

    extracted = extract_text(f, max_chars=2000)
    assert len(extracted) == 2000


def test_extract_cap_file_size(tmp_path):
    f = tmp_path / "huge.txt"
    f.write_text("x" * 200)

    # Exceed artificial small max_file_size
    extracted = extract_text(f, max_file_size=100)
    assert extracted == ""


def test_extract_nonexistent_file(tmp_path):
    assert extract_text(tmp_path / "does_not_exist.txt") == ""
